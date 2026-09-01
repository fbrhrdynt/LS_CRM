"""Excel import/export using openpyxl."""
import io
from typing import Callable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill(start_color="111111", end_color="111111", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=10)


def _autosize(ws):
    for col in ws.columns:
        max_len = 0
        letter = get_column_letter(col[0].column)
        for cell in col:
            v = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(v))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 45)


def rows_to_xlsx(sheet_name: str, headers: list, rows: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    ws.append(headers)
    for c in ws[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="left", vertical="center")
    for r in rows:
        ws.append(r)
    _autosize(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_to_dicts(content: bytes) -> list[dict]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h else "" for h in rows[0]]
    out = []
    for row in rows[1:]:
        if row is None or all(v is None for v in row):
            continue
        d = {}
        for i, h in enumerate(headers):
            if not h:
                continue
            d[h] = row[i] if i < len(row) else None
        out.append(d)
    return out
